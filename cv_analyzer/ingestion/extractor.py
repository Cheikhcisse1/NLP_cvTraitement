"""
ingestion/extractor.py
Extraction de texte brut à partir de PDF, DOCX et TXT.
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Union

from loguru import logger


def extract_text(source: Union[str, Path, bytes], filename: str = "") -> str:
    """
    Point d'entrée unique.
    Accepte un chemin fichier OU des bytes (upload HTTP).
    Renvoie le texte brut extrait.
    """
    if isinstance(source, (str, Path)):
        path = Path(source)
        suffix = path.suffix.lower()
        with open(path, "rb") as f:
            data = f.read()
    else:
        data = source
        suffix = Path(filename).suffix.lower() if filename else _guess_type(data)

    logger.debug(f"Extraction | type={suffix} | taille={len(data)} octets")

    if suffix == ".pdf":
        return _from_pdf(data)
    elif suffix in (".docx", ".doc"):
        return _from_docx(data)
    elif suffix in (".txt", ".md", ""):
        return data.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Format non supporté : {suffix!r}")


# ── Extracteurs internes ───────────────────────────────────────────────────────

def _from_pdf(data: bytes) -> str:
    """Extrait le texte page par page avec pdfplumber."""
    import pdfplumber

    text_parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            t = page.extract_text(x_tolerance=2, y_tolerance=2)
            if t:
                text_parts.append(t)
    return "\n\n".join(text_parts)


def _from_docx(data: bytes) -> str:
    """Extrait le texte paragraphe par paragraphe depuis un DOCX."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Inclut aussi les tableaux (compétences souvent en tableau)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    return "\n".join(paragraphs)


def _guess_type(data: bytes) -> str:
    """Devine le type fichier depuis la magic signature."""
    if data[:4] == b"%PDF":
        return ".pdf"
    if data[:2] == b"PK":          # ZIP → DOCX
        return ".docx"
    return ".txt"
